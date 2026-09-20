from pathlib import Path
from openpyxl import load_workbook
from langchain_core.documents import Document


def load_excel(file_path):
    path = Path(file_path)

    # load_workbook 是 openpyxl 提供的方法，用来读取 .xlsx Excel 文件
    # data_only=True 表示如果单元格中有公式，优先读取公式计算后的值
    workbook = load_workbook(
        filename=path,
        data_only=True
    )

    documents = []

    # workbook.worksheets 获取 Excel 中的所有 Sheet
    for sheet in workbook.worksheets:

        # iter_rows(values_only=True) 按行读取单元格
        # values_only=True 表示直接获取单元格的值，而不是 Cell 对象
        rows = sheet.iter_rows(values_only=True)

        # 第一行作为表头
        headers = next(rows)

        # 从第 2 行开始读取真正的数据
        for row_number, row in enumerate(rows, start=2):

            data = {}

            # 将表头和每一行数据组合成字典
            for header, value in zip(headers, row):

                if header is not None:
                    data[str(header)] = value

            # 跳过完全为空的行
            if not any(value is not None for value in row):
                continue

            # 将一行 Excel 数据转换成文本
            content_parts = []

            for key, value in data.items():

                if value is not None:
                    content_parts.append(f"{key}: {value}")

            content = "\n".join(content_parts)

            # 每一行 Excel 数据转换成一个 Document
            document = Document(
                page_content=content,
                metadata={
                    "source": path.name,
                    "suffix": path.suffix,
                    "sheet": sheet.title,
                    "row": row_number
                }
            )

            documents.append(document)

    # close 关闭 Excel 文件，释放资源
    workbook.close()

    return documents


# 当前 Python 文件所在目录
BASE_DIR = Path(__file__).parent

# 加载 Excel
documents = load_excel(
    BASE_DIR / "data" / "rag_excel_test.xlsx"
)

print("Document 数量:", len(documents))

for index, document in enumerate(documents, start=1):
    print()
    print(f"Document {index}")
    print("content:")
    print(document.page_content)
    print("metadata:")
    print(document.metadata)