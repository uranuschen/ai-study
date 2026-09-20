import json
from pathlib import Path
from agent_learning.core.models import Document


# 方法功能：递归展开嵌套 JSON，保留对象路径和数组下标，让业务字段可以参与向量化。
# 输入：JSON 值与可选字段路径；输出：可读的多行文本。
# 把嵌套 JSON 转成保留字段路径的文本。
def flatten_json(value, prefix=""):
    # 按字段名排序，保证重复导入内容稳定。
    if isinstance(value, dict):
        # 展开对象。
        return "\n".join(flatten_json(value[key], f"{prefix}.{key}".strip(".")) for key in sorted(value))
    # 保留数组元素位置。
    if isinstance(value, list):
        # 展开数组。
        return "\n".join(flatten_json(item, f"{prefix}[{index}]") for index, item in enumerate(value))
    # 保留布尔值、数值和空值含义。
    return f"{prefix}: {json.dumps(value, ensure_ascii=False)}"


# 类职责：读取不同格式资料，将 JSON、文本、PDF、Word、Excel 转换为统一 Document 列表。
# 所有文件格式统一输出 Document。
class DocumentLoader:
    # 方法功能：根据文件后缀选择加载方式，并保留页码、表名或行号等来源。
    # 输入：文件 Path 和默认租户；输出：Document 列表；不支持的格式会报错。
    # 只读取用户指定文件。
    def load(self, path: Path, tenant_id="acme") -> list[Document]:
        # 标准化后缀。
        suffix = path.suffix.lower()
        # 企业结构化数据入口。
        if suffix in {".json", ".jsonl"}:
            # 使用 UTF-8。
            text = path.read_text(encoding="utf-8")
            # 解析记录。
            records = [json.loads(line) for line in text.splitlines() if line.strip()] if suffix == ".jsonl" else json.loads(text)
            # 兼容单对象。
            records = records if isinstance(records, list) else [records]
            # 转换每条记录。
            return [self.from_record(record, str(path), tenant_id) for record in records]
        # 普通文本保留完整原文。
        if suffix in {".md", ".txt"}:
            # 输出一篇文档。
            return [self._plain(path, path.read_text(encoding="utf-8"), tenant_id, {})]
        # PDF 每页独立保留页码。
        if suffix == ".pdf":
            from pypdf import PdfReader
            # 解析失败直接报错，避免静默漏页。
            reader = PdfReader(path)
            # 空页在校验时提示 OCR。
            return [self._plain(path, page.extract_text() or "", tenant_id, {"page": index}, f":page:{index}") for index, page in enumerate(reader.pages, 1)]
        # 同时读取 Word 段落和表格。
        if suffix == ".docx":
            from docx import Document as WordDocument
            # 加载 Word。
            doc = WordDocument(path)
            # 提取段落。
            parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
            # 提取表格行。
            parts.extend(" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
            # 合并为一篇文档。
            return [self._plain(path, "\n\n".join(parts), tenant_id, {})]
        # Excel 按 Sheet 和行记录。
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            # 使用公式缓存值。
            workbook = load_workbook(path, read_only=True, data_only=True)
            # 收集非空记录。
            documents = []
            # 保证文件句柄释放。
            try:
                # 遍历工作表。
                for sheet in workbook:
                    # 流式读取单元格。
                    rows = sheet.iter_rows(values_only=True)
                    # 首行作为表头。
                    headers = next(rows, ())
                    # 保留 Excel 行号。
                    for index, row in enumerate(rows, 2):
                        # 跳过完全空行。
                        if not any(value is not None for value in row):
                            # 不制造空文档。
                            continue
                        # 保留字段名。
                        content = "\n".join(f"{header or f'列{column}'}: {value}" for column, (header, value) in enumerate(zip(headers, row), 1) if value is not None)
                        # 定位到行。
                        documents.append(self._plain(path, content, tenant_id, {"sheet": sheet.title, "row": index}, f":{sheet.title}:{index}"))
            # 无论成功与否都关闭文件。
            finally:
                # 释放句柄。
                workbook.close()
            # 返回行文档。
            return documents
        # 明确告知不支持的输入。
        raise ValueError(f"不支持的格式: {suffix}")

    # 方法功能：按 vector_fields 投影业务 JSON，同时保留完整 payload 与元数据。
    # 输入：一条 JSON 对象、来源和默认租户；输出：一个 Document。
    # 将业务 JSON 映射为文档。
    def from_record(self, record: dict, source: str, tenant_id="acme") -> Document:
        # 根记录必须为对象。
        if not isinstance(record, dict):
            # 防止错误结构。
            raise ValueError("JSON 每条记录必须是对象")
        # 明确选择哪些字段参与向量化。
        fields = record.get("vector_fields", ["content", "structured_data"])
        # 校验白名单。
        if not isinstance(fields, list) or not fields or not all(isinstance(key, str) for key in fields):
            # 避免意外向量化整个对象。
            raise ValueError("vector_fields 必须是非空字符串数组")
        # 标题提供语义上下文。
        parts = [str(record.get("title", ""))]
        # 仅向量化指定顶层字段。
        for key in fields:
            # 默认字段可以缺失。
            if key in record:
                # 结构化值可读化。
                parts.append(record[key] if isinstance(record[key], str) else flatten_json(record[key], key))
        # 形成稳定的向量输入。
        content = "\n\n".join(part for part in parts if part.strip())
        # 保留原始元数据。
        metadata = dict(record.get("metadata", {}))
        # 记录投影配置以便追溯。
        metadata["vector_fields"] = fields
        # 受支持的业务字段。
        names = ["external_id", "title", "business_module", "doc_type", "security_level", "version", "is_latest", "tags"]
        # 忽略额外字段但在 payload 保留。
        values = {key: record[key] for key in names if key in record}
        # 返回统一文档。
        return Document(**values, content=content, tenant_id=record.get("tenant_id", tenant_id), source=record.get("source", source), metadata=metadata, payload=record)

    # 方法功能：为非 JSON 文件补齐统一文档字段和稳定来源标识。
    # 输入：路径、正文、租户、元数据及可选位置后缀；输出：一个 Document。
    # 为非 JSON 文件构造通用结构。
    def _plain(self, path, content, tenant_id, metadata, suffix=""):
        # 路径构成稳定标识。
        return Document(external_id=f"{path.as_posix()}{suffix}", title=path.stem, content=content, tenant_id=tenant_id, source=str(path), metadata=metadata)
