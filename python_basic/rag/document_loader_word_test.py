from pathlib import Path
from docx import Document as WordDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter


def load_word(file_path):
    path = Path(file_path)

    # 读取 Word 文件
    word_doc = WordDocument(path)

    paragraphs = []

    # 提取 Word 中的非空段落
    # strip 去掉字符串开头和结尾的空白字符
    for paragraph in word_doc.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    # 将所有段落拼成统一文本
    content = "\n\n".join(paragraphs)

    # 转成 LangChain 统一的 Document 结构
    document = Document(
        page_content=content,
        metadata={
            "source": path.name,
            "suffix": path.suffix
        }
    )

    return document


# 当前 Python 文件所在目录
BASE_DIR = Path(__file__).parent

# 加载 Word
document = load_word(BASE_DIR / "data" / "refund_test.docx")

print("原始 Document")
print("content:")
print(document.page_content)
print("metadata:")
print(document.metadata)

# 创建文本切分器
text_splitter = RecursiveCharacterTextSplitter(
    chunk_size=50,
    chunk_overlap=10
)

# Document 可以直接进行切块
chunks = text_splitter.split_documents([document])

print()
print("Chunk 数量:", len(chunks))

for index, chunk in enumerate(chunks, start=1):
    print()
    print(f"Chunk {index}")
    print("content:")
    print(chunk.page_content)
    print("metadata:")
    print(chunk.metadata)