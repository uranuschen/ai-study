from pathlib import Path
import json

from pypdf import PdfReader
from docx import Document as WordDocument
from openpyxl import load_workbook

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from sentence_transformers import SentenceTransformer
import psycopg2


# 当前 Python 文件所在目录
BASE_DIR = Path(__file__).parent

# 需要扫描的数据目录
DATA_DIR = BASE_DIR / "data"


# PostgreSQL 配置
DB_CONFIG = {
    "host": "localhost",
    "port": 5432,
    "database": "ragdb",
    "user": "rag",
    "password": "rag123456"
}


# 支持的文件类型
SUPPORTED_SUFFIXES = {
    ".pdf",
    ".md",
    ".docx",
    ".xlsx"
}


# 加载 PDF
def load_pdf(file_path):
    path = Path(file_path)

    # PdfReader 是 pypdf 提供的方法，用来读取 PDF
    reader = PdfReader(path)

    documents = []

    # reader.pages 获取 PDF 中的所有页面
    for page_number, page in enumerate(reader.pages, start=1):

        try:
            # extract_text() 提取当前 PDF 页面的文字
            content = page.extract_text()

        except Exception as e:
            # 某一页解析失败时跳过，不影响其他文件继续入库
            print(
                f"PDF 页面解析失败: "
                f"{path.name} 第 {page_number} 页, "
                f"错误: {e}"
            )
            continue

        if content:
            content = content.strip()

        # 空白页直接跳过
        if not content:
            continue

        # PDF 每一页生成一个 Document
        document = Document(
            page_content=content,
            metadata={
                "source": path.name,
                "suffix": path.suffix,
                "page": page_number
            }
        )

        documents.append(document)

    return documents


# 加载 Markdown
def load_markdown(file_path):
    path = Path(file_path)

    # read_text() 是 pathlib 提供的方法，用来直接读取文本文件
    content = path.read_text(encoding="utf-8")

    document = Document(
        page_content=content,
        metadata={
            "source": path.name,
            "suffix": path.suffix
        }
    )

    return [document]


# 加载 Word
def load_word(file_path):
    path = Path(file_path)

    # WordDocument 来自 python-docx，用来读取 .docx 文件
    word_doc = WordDocument(path)

    paragraphs = []

    # word_doc.paragraphs 获取 Word 中的所有段落
    for paragraph in word_doc.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    # 将多个段落拼成一个完整文本
    content = "\n\n".join(paragraphs)

    document = Document(
        page_content=content,
        metadata={
            "source": path.name,
            "suffix": path.suffix
        }
    )

    return [document]


# 加载 Excel
def load_excel(file_path):
    path = Path(file_path)

    # load_workbook 是 openpyxl 提供的方法，用来读取 .xlsx
    # data_only=True 优先读取公式计算后的值
    workbook = load_workbook(
        filename=path,
        data_only=True
    )

    documents = []

    # workbook.worksheets 获取所有 Sheet
    for sheet in workbook.worksheets:

        # iter_rows(values_only=True) 按行读取，并直接返回单元格的值
        rows = sheet.iter_rows(values_only=True)

        try:
            # 第一行作为表头
            headers = next(rows)
        except StopIteration:
            continue

        # 从 Excel 第 2 行开始读取数据
        for row_number, row in enumerate(rows, start=2):

            # 跳过完全为空的行
            if not any(value is not None for value in row):
                continue

            content_parts = []

            # zip() 将表头和当前行的数据一一对应
            for header, value in zip(headers, row):

                if header is not None and value is not None:
                    content_parts.append(
                        f"{header}: {value}"
                    )

            content = "\n".join(content_parts)

            if not content:
                continue

            # Excel 每一行生成一个 Document
            document = Document(
                page_content=content,
                metadata={
                    "source": path.name,
                    "suffix": path.suffix,
                    "sheet": sheet.title,
                    "row": row_number
                }
            )

            documents.append(document)

    # 关闭 Excel 文件
    workbook.close()

    return documents


# 根据文件后缀自动选择 Loader
def load_document(file_path):
    path = Path(file_path)

    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return load_pdf(path)

    elif suffix == ".md":
        return load_markdown(path)

    elif suffix == ".docx":
        return load_word(path)

    elif suffix == ".xlsx":
        return load_excel(path)

    else:
        raise ValueError(
            f"暂不支持的文件类型: {suffix}"
        )


# 扫描 data 目录中的所有支持文件
def load_all_documents(data_dir):
    documents = []

    # iterdir() 是 pathlib 提供的方法，用来遍历目录中的文件
    for file_path in data_dir.iterdir():

        if not file_path.is_file():
            continue

        suffix = file_path.suffix.lower()

        if suffix not in SUPPORTED_SUFFIXES:
            continue

        print()
        print("正在加载:", file_path.name)

        try:
            file_documents = load_document(file_path)

            print(
                "生成 Document 数量:",
                len(file_documents)
            )

            documents.extend(file_documents)
        except Exception as e:
            # 单个文件失败时跳过，继续处理其他文件
            print(
                f"文件加载失败: {file_path.name}"
            )
            print(
                f"错误信息: {e}"
            )

    return documents


# 对所有 Document 统一进行切块
def split_all_documents(documents):

    # RecursiveCharacterTextSplitter 是 LangChain 的递归文本切分器
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=200,
        chunk_overlap=30
    )

    chunks = []

    for document in documents:

        # split_documents() 接收 Document 列表
        # 切块后会自动保留原来的 metadata
        document_chunks = text_splitter.split_documents(
            [document]
        )

        # 给同一个 Document 的 Chunk 增加编号
        for chunk_index, chunk in enumerate(
            document_chunks,
            start=1
        ):

            chunk.metadata["chunk_index"] = chunk_index

            chunks.append(chunk)

    return chunks


# 初始化 PostgreSQL 表
def init_database(conn):

    cursor = conn.cursor()

    # 创建 pgvector 扩展
    cursor.execute(
        """
        CREATE EXTENSION IF NOT EXISTS vector
        """
    )

    # 第一次运行时创建表
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS rag_documents (
            id BIGSERIAL PRIMARY KEY,
            content TEXT NOT NULL,
            metadata JSONB,
            source VARCHAR(255),
            embedding VECTOR(512)
        )
        """
    )

    # ALTER TABLE 用来修改已经存在的旧表
    # metadata 不存在时增加
    cursor.execute(
        """
        ALTER TABLE rag_documents
        ADD COLUMN IF NOT EXISTS metadata JSONB
        """
    )

    # source 不存在时增加
    cursor.execute(
        """
        ALTER TABLE rag_documents
        ADD COLUMN IF NOT EXISTS source VARCHAR(255)
        """
    )

    conn.commit()

    cursor.close()


# 删除以前已经入库的相同文件
def delete_old_documents(conn, source_names):

    cursor = conn.cursor()

    for source in source_names:

        cursor.execute(
            """
            DELETE FROM rag_documents
            WHERE source = %s
            """,
            (source,)
        )

    conn.commit()

    cursor.close()


# Embedding 并写入数据库
def insert_documents(conn, chunks):

    # SentenceTransformer 是 sentence-transformers 提供的模型加载类
    model = SentenceTransformer(
        "BAAI/bge-small-zh-v1.5"
    )

    # 提取所有 Chunk 文本
    texts = [
        chunk.page_content
        for chunk in chunks
    ]

    print()
    print("开始生成 Embedding...")

    # encode() 将多段文本一次性转换成向量
    # normalize_embeddings=True 将向量归一化，方便后续计算余弦相似度
    embeddings = model.encode(
        texts,
        normalize_embeddings=True
    )

    cursor = conn.cursor()

    for chunk, embedding in zip(
        chunks,
        embeddings
    ):

        # numpy ndarray 转成 Python list
        embedding_list = embedding.tolist()

        # pgvector 可以接收这种 [0.1,0.2,...] 格式
        embedding_string = str(embedding_list)

        # json.dumps() 将 Python dict 转成 JSON 字符串
        metadata_json = json.dumps(
            chunk.metadata,
            ensure_ascii=False
        )

        cursor.execute(
            """
            INSERT INTO rag_documents (
                content,
                metadata,
                source,
                embedding
            )
            VALUES (
                %s,
                %s::jsonb,
                %s,
                %s::vector
            )
            """,
            (
                chunk.page_content,
                metadata_json,
                chunk.metadata.get("source"),
                embedding_string
            )
        )

    conn.commit()

    cursor.close()


def main():

    print("开始扫描目录:")
    print(DATA_DIR)

    # 第一步：加载所有文件
    documents = load_all_documents(
        DATA_DIR
    )

    print()
    print("原始 Document 总数:", len(documents))

    # 第二步：统一切 Chunk
    chunks = split_all_documents(
        documents
    )

    print("Chunk 总数:", len(chunks))

    print()
    print("Chunk 示例:")

    for index, chunk in enumerate(
        chunks[:5],
        start=1
    ):
        print()
        print(f"Chunk {index}")
        print(chunk.page_content)
        print(chunk.metadata)

    # psycopg2.connect() 创建 PostgreSQL 数据库连接
    conn = psycopg2.connect(
        **DB_CONFIG
    )

    try:

        # 第三步：初始化数据库
        init_database(conn)

        # 获取本次所有文件名
        source_names = {
            chunk.metadata["source"]
            for chunk in chunks
        }

        # 删除这些文件以前的旧数据，避免重复入库
        delete_old_documents(
            conn,
            source_names
        )

        # 第四步：Embedding + 入库
        insert_documents(
            conn,
            chunks
        )

        print()
        print("全部入库完成")
        print("Chunk 数量:", len(chunks))

    finally:

        # 关闭数据库连接
        conn.close()


if __name__ == "__main__":
    main()